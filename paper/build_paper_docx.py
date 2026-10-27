"""
build_paper_docx.py -- Word copy of the IEEE Access manuscript (paper/latex/main.tex), for
readers and similarity checks that need a .docx. Laid out like the IEEE PDF: two columns,
wide figures and tables across both. The LaTeX PDF stays the version of record.

Nothing is retyped. The LaTeX source is converted by pandoc after these substitutions:
  \\val{id}      value from paper/latex/generated/values.tex (the build fails on an unknown id)
  \\cite{..}     IEEE number from the order of paper/latex/build/main.bbl
  \\ref{..}      section, figure, table, algorithm or equation number, counted as LaTeX does
  figures       paper/figures/<name>.png in place of generated/figures/<name>.pdf
  references    as IEEEtran.bst wrote them to paper/latex/build/main.bbl

Needs pandoc: pip install pypandoc_binary. Run python paper/build_latex.py first, so that
generated/ and build/main.bbl are current.

Output: paper/latex/build/manuscript_ieee.docx
Run: python paper/build_paper_docx.py
"""
from __future__ import annotations

import copy
import os
import re
import sys

import pypandoc
from lxml import etree
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_docx import add_runs

HERE = os.path.dirname(os.path.abspath(__file__))
LATEX = os.path.join(HERE, "latex")
OUT = os.path.join(LATEX, "build", "manuscript_ieee.docx")
FONT = "Times New Roman"
ROMAN = "I II III IV V VI VII VIII IX X XI XII XIII XIV XV".split()
REF_MARK = "REFERENCE-LIST-GOES-HERE"
ALG_MARK = "ALGLINE%sQQ"                    # finish() styles and removes these markers
WIDE_START, WIDE_END, TITLE_END = "WIDESTARTQQ", "WIDEENDQQ", "TITLEENDQQ"
# IEEE two-column page: US Letter, margins 0.75 in top, 1 in bottom, 0.625 in sides,
# 0.2 in between columns (the IEEE Word template's values).
TEXT_W_CM = 2.54 * (8.5 - 2 * 0.625)
COL_W_CM = (TEXT_W_CM - 2.54 * 0.2) / 2


def read(path):
    with open(path, encoding="utf-8") as fh:
        return fh.read()


def strip_comments(text):
    """Drop comments; a line holding only a comment goes too, so it cannot split a paragraph."""
    return "\n".join(re.sub(r"(?<!\\)%.*", "", line) for line in text.splitlines()
                     if not line.lstrip().startswith("%"))


def group_end(text, i):
    """Index just past the brace group that opens at text[i] == '{'."""
    depth = 0
    for j in range(i, len(text)):
        if text[j] == "{" and text[j - 1] != "\\":
            depth += 1
        elif text[j] == "}" and text[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return j + 1
    raise ValueError("unbalanced braces after %r" % text[i:i + 60])


def replace_command(text, name, fn, nargs=1):
    """Replace \\name[opt]{a}{b}.. (balanced braces) by fn(a, b, ..)."""
    pat = re.compile(r"\\%s(\[[^\]]*\])?(?=\{)" % name)
    while True:
        m = pat.search(text)
        if not m:
            return text
        args, j = [], m.end()
        for _ in range(nargs):
            k = group_end(text, j)
            args.append(text[j + 1:k - 1])
            j = k
        text = text[:m.start()] + fn(*args) + text[j:]


def unwrap_table(m):
    """A table float becomes its caption line followed by the bare tabular, which pandoc
    converts reliably; captions inside floats and \\resizebox wrappers are otherwise lost."""
    body = m.group(0)
    i = body.index("\\caption{") + len("\\caption")
    caption = body[i + 1:group_end(body, i) - 1]
    tabular = re.search(r"\\begin\{tabular\}.*?\\end\{tabular\}", body, re.S).group(0)
    out = "\n\n\\noindent %s\n\n%s\n\n" % (caption, tabular)
    if body.startswith("\\begin{table*}"):
        out = "\n\n%s\n\n%s\n\n%s\n\n" % (WIDE_START, out, WIDE_END)
    return out


def expand_inputs(text):
    def inc(m):
        name = m.group(1)
        if name.endswith("values"):
            return ""
        path = os.path.join(LATEX, name if name.endswith(".tex") else name + ".tex")
        return expand_inputs(strip_comments(read(path)))
    return re.sub(r"\\input\{([^}]+)\}", inc, text)


def load_values():
    vals = {}
    for m in re.finditer(r"\\@namedef\{val@([a-z0-9_]+)\}\{(.*)\}%",
                         read(os.path.join(LATEX, "generated", "values.tex"))):
        vals[m.group(1)] = m.group(2).replace("\\ensuremath{-}", "\u2212")
    return vals


def bbl_order():
    return re.findall(r"\\bibitem\{([^}]+)\}", read(os.path.join(LATEX, "build", "main.bbl")))


def compress(nums):
    """IEEE style, as the cite package prints it: [1], [3]–[5]."""
    nums, out, i = sorted(set(nums)), [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append("[%d]\u2013[%d]" % (nums[i], nums[j]) if j - i >= 2 else
                   ", ".join("[%d]" % n for n in nums[i:j + 1]))
        i = j + 1
    return ", ".join(out)


def number_labels(text):
    """Assign each \\label the number LaTeX would print for it, and number the headings,
    captions and equations in the text, since pandoc numbers none of them."""
    labels, counts, current = {}, {"sec": 0, "sub": 0, "fig": 0, "tab": 0, "alg": 0, "eq": 0}, None
    event = re.compile(r"\\(section|subsection)(\*?)\{|\\begin\{(figure|table|algorithm|equation)"
                       r"\*?\}|\\caption\{|\\label\{([^}]+)\}|\\appendices")
    out, pos, env, appendix = [], 0, None, False
    for m in event.finditer(text):
        out.append(text[pos:m.start()])
        pos = m.end()
        kind, star, envname, label = m.group(1), m.group(2), m.group(3), m.group(4)
        if m.group(0) == "\\appendices":
            appendix, counts["sec"] = True, 0
        elif kind == "section" and appendix:
            counts["sec"] += 1
            current = ("app", chr(64 + counts["sec"]))
            out.append("\\section{Appendix %s: " % current[1])
        elif kind == "section" and not star:
            counts["sec"] += 1
            counts["sub"] = 0
            current = ("sec", ROMAN[counts["sec"] - 1])
            out.append("\\section{%s. " % current[1])
        elif kind == "subsection" and not star:
            counts["sub"] += 1
            current = ("sub", "%s-%s" % (ROMAN[counts["sec"] - 1], chr(64 + counts["sub"])))
            out.append("\\subsection{%s. " % chr(64 + counts["sub"]))
        elif kind:
            out.append(m.group(0))
        elif envname:
            key = {"figure": "fig", "table": "tab", "algorithm": "alg", "equation": "eq"}[envname]
            counts[key] += 1
            env = key
            n = counts[key]
            current = (key, ROMAN[n - 1] if key == "tab" else str(n))
            out.append(m.group(0))
        elif label is None:                                   # \caption{
            prefix = {"fig": "Fig. %s. ", "tab": "TABLE %s. ", "alg": "Algorithm %s: "}
            out.append("\\caption{" + (prefix[env] % current[1] if env in prefix else ""))
        else:
            labels[label] = current[1]
    out.append(text[pos:])
    return "".join(out), labels


def convert_algorithm(m):
    """Algorithm 1 line by line as algpseudocode prints it: numbered lines, bold keywords
    (if/then/else/end if/return), comments after a triangle. Each line is a paragraph with a
    marker that finish() turns into formatting."""
    body = m.group(0)
    i = body.index("\\caption{") + len("\\caption")
    cap = body[i + 1:group_end(body, i) - 1]
    num, cap = re.match(r"Algorithm (\d+): (.*)", cap, re.S).groups()
    steps = re.search(r"\\begin\{algorithmic\}(?:\[\d\])?(.*?)\\end\{algorithmic\}", body,
                      re.S).group(1)
    out = ["%s\\textbf{Algorithm %s} %s" % (ALG_MARK % "cap", num, cap.strip())]
    depth, n = 0, 0
    for cmd, rest in re.findall(r"\\(Require|State|If|Else|EndIf)(.*?)"
                                r"(?=\\Require|\\State|\\If|\\Else|\\EndIf|$)", steps, re.S):
        rest = " ".join(rest.split())
        rest = re.sub(r"\\Comment\{(.*?)\}$", "\u22b2 \\1", rest)
        rest = rest.replace("\\Return", "\\textbf{return}")
        if cmd == "Require":
            out.append("%s\\textbf{Input:} %s" % (ALG_MARK % "in", rest))
            continue
        if cmd in ("Else", "EndIf"):
            depth -= 1
        n += 1
        if cmd == "If":
            cond = rest[1:group_end(rest, 0) - 1]
            text = "\\textbf{if} %s \\textbf{then}" % cond
        else:
            text = {"State": rest, "Else": "\\textbf{else}", "EndIf": "\\textbf{end if}"}[cmd]
        out.append("%s%d: %s" % (ALG_MARK % depth, n, text))
        if cmd in ("If", "Else"):
            depth += 1
    out.append(ALG_MARK % "end")
    return "\n\n" + "\n\n".join(out) + "\n\n"


def to_latex():
    vals, order = load_values(), bbl_order()
    tex = expand_inputs(strip_comments(read(os.path.join(LATEX, "main.tex"))))
    tex = re.sub(r"\\makeatletter\s*\\newcommand\{\\val\}.*?\\makeatother", "", tex, flags=re.S)

    missing = []

    def val(m):
        if m.group(1) not in vals:
            missing.append(m.group(1))
            return "[SOURCE NOT FOUND: %s]" % m.group(1)
        return vals[m.group(1)]
    tex = re.sub(r"\\val\{([a-z0-9_]+)\}", val, tex)

    # \thanks notes under the title: pandoc drops them, so finish() adds them as footnotes.
    notes = []
    tex = replace_command(tex, "thanks",
                          lambda a: notes.append(" ".join(a.replace("~", " ").split())) or "")

    def cite(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        unknown = [k for k in keys if k not in order]
        if unknown:
            missing.extend("cite:" + k for k in unknown)
        return compress([order.index(k) + 1 for k in keys if k in order])
    tex = re.sub(r"\\cite\{([^}]+)\}", cite, tex)

    tex, labels = number_labels(tex)

    def ref(m):
        if m.group(1) not in labels:
            missing.append("ref:" + m.group(1))
            return "??"
        return labels[m.group(1)]
    tex = re.sub(r"\\ref\{([^}]+)\}", ref, tex)
    tex = re.sub(r"\\label\{[^}]+\}", "", tex)
    tex = re.sub(r"\\begin\{algorithm\}.*?\\end\{algorithm\}", convert_algorithm, tex, flags=re.S)

    eq = iter(range(1, 100))
    tex = re.sub(r"\\begin\{equation\}(.*?)\\end\{equation\}",
                 lambda m: "\\[%s \\qquad (%d)\\]" % (m.group(1).strip(), next(eq)), tex, flags=re.S)

    def graphic(m):
        name = os.path.splitext(os.path.basename(m.group(2)))[0]
        png = os.path.join(HERE, "figures", name + ".png")
        if not os.path.isfile(png):
            missing.append("figure:" + name)
        w = re.search(r"width=([\d.]*)\\(textwidth|columnwidth)", m.group(1))
        width = float(w.group(1) or 1) * (TEXT_W_CM if w.group(2) == "textwidth" else COL_W_CM)
        return "\\includegraphics[width=%.2fcm]{%s}" % (width, png.replace("\\", "/"))
    tex = re.sub(r"\\includegraphics\[([^\]]*)\]\{([^}]+)\}", graphic, tex)
    tex = re.sub(r"\\begin\{figure\*\}(.*?)\\end\{figure\*\}",
                 lambda m: "\n\n%s\n\n\\begin{figure}%s\\end{figure}\n\n%s\n\n" % (
                     WIDE_START, m.group(1), WIDE_END), tex, flags=re.S)

    tex = replace_command(tex, "makecell", lambda a: a.replace("\\\\", " "))
    tex = replace_command(tex, "IEEEPARstart", lambda a, b: a + b, nargs=2)
    tex = re.sub(r"\\begin\{table\*?\}.*?\\end\{table\*?\}", unwrap_table, tex, flags=re.S)
    tex = re.sub(r"\\(scriptsize|footnotesize|small)\b|\\setlength\{[^}]*\}\{[^}]*\}", "", tex)
    tex = tex.replace("\\begin{IEEEkeywords}", "\\textit{Index Terms}---").replace(
        "\\end{IEEEkeywords}", "\n\n%s\n\n" % TITLE_END)
    tex = re.sub(r"\\begin\{IEEEbiographynophoto\}\{([^}]*)\}", r"\n\n\\noindent\\textbf{\1} ", tex)
    tex = tex.replace("\\end{IEEEbiographynophoto}", "\n\n")
    tex = re.sub(r"\\markboth\{.*?\}\{.*?\}\s*\n", "\n", tex, flags=re.S)
    tex = re.sub(r"\\bibliographystyle\{[^}]*\}", "", tex)
    tex = re.sub(r"\\bibliography\{[^}]*\}", "\n\n%s\n\n" % REF_MARK, tex)
    if missing:
        sys.exit("unresolved: " + ", ".join(missing))
    return tex, order, notes


def rule(parent_tag, sides, size=4):
    """A border element (tblBorders, tcBorders or pBdr) with single lines on the given sides."""
    el = OxmlElement(parent_tag)
    for side in sides:
        s = OxmlElement("w:" + side)
        s.set(qn("w:val"), "single")
        s.set(qn("w:sz"), str(size))
        s.set(qn("w:space"), "1")
        s.set(qn("w:color"), "000000")
        el.append(s)
    return el


def booktabs(table):
    """Top and bottom rules and a rule under the header row, no vertical lines, full width."""
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblBorders", "w:tblW", "w:tblStyle", "w:tblLook"):
        for old in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(old)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    tbl_pr.append(width)
    tbl_pr.append(rule("w:tblBorders", ("top", "bottom"), 8))
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_pr.append(rule("w:tcBorders", ("bottom",), 4))


def bbl_entries():
    """Reference entries exactly as IEEEtran.bst wrote them to build/main.bbl, converted to
    plain text with *italics*."""
    text = read(os.path.join(LATEX, "build", "main.bbl"))
    entries = []
    for m in re.finditer(r"\\bibitem\{[^}]+\}(.*?)(?=\\bibitem|\\end\{thebibliography\})", text,
                         re.S):
        e = re.sub(r"\\BIBentry(ALT|STD)interwordspacing", "", m.group(1))
        e = replace_command(e, "emph", lambda a: "*%s*" % a)
        e = replace_command(e, "url", lambda a: a)
        e = replace_command(e, "BIBforeignlanguage", lambda a, b: b, nargs=2)
        e = e.replace("\\newblock", " ").replace("``", "\u201c").replace("''", "\u201d")
        e = e.replace("---", "\u2014").replace("--", "\u2013").replace("~", "\u00a0")
        e = re.sub(r"\\([&_%#$])", r"\1", e).replace("{", "").replace("}", "")
        entries.append(" ".join(e.split()))
    return entries


def section_break(par, cols, template):
    """Turn a marker paragraph into a continuous section break that ends a section of
    `cols` columns."""
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    sect = copy.deepcopy(template)
    for old in sect.findall(qn("w:headerReference")) + sect.findall(qn("w:footerReference")):
        sect.remove(old)
    set_columns(sect, cols)
    ppr = par._p.get_or_add_pPr()
    ppr.append(sect)
    par.paragraph_format.space_before = par.paragraph_format.space_after = Pt(0)


def set_columns(sect, cols):
    for tag in ("w:type", "w:cols"):
        for old in sect.findall(qn(tag)):
            sect.remove(old)
    kind = OxmlElement("w:type")
    kind.set(qn("w:val"), "continuous")
    sect.insert(0, kind)
    col = OxmlElement("w:cols")
    col.set(qn("w:num"), str(cols))
    col.set(qn("w:space"), str(int(0.2 * 1440)))
    sect.insert(list(sect).index(sect.find(qn("w:pgMar"))) + 1, col)


def strip_marker(par, marker):
    for r in par.runs:
        if marker in r.text:
            r.text = r.text.replace(marker, "", 1)
            return


def author_notes(doc, notes):
    """The \\thanks notes as unnumbered footnotes on the author line, as IEEEtran prints
    them at the foot of the first column."""
    part = next(r.target_part for r in doc.part.rels.values() if r.reltype.endswith("/footnotes"))
    root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
    nid = 1 + max(int(f.get(qn("w:id"))) for f in root.findall(qn("w:footnote")))
    author = next(p for p in doc.paragraphs if p.style.name == "Author")
    for text in notes:
        fn = etree.SubElement(root, qn("w:footnote"))
        fn.set(qn("w:id"), str(nid))
        p = etree.SubElement(fn, qn("w:p"))
        etree.SubElement(etree.SubElement(p, qn("w:pPr")), qn("w:pStyle")).set(
            qn("w:val"), "FootnoteText")
        t = etree.SubElement(etree.SubElement(p, qn("w:r")), qn("w:t"))
        t.text = text
        ref = OxmlElement("w:r")
        mark = OxmlElement("w:footnoteReference")
        mark.set(qn("w:customMarkFollows"), "1")
        mark.set(qn("w:id"), str(nid))
        ref.append(mark)
        author._p.append(ref)
        nid += 1
    if not hasattr(part, "_element"):
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                    standalone=True)


def finish(path, order, notes):
    doc = Document(path)
    author_notes(doc, notes)
    for st in doc.styles:
        if st.type != 1:
            continue
        st.font.name = FONT
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(a), None)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
        if st.name.startswith(("Heading", "Title", "Subtitle")):
            st.font.color.rgb = RGBColor(0, 0, 0)
    names = [s.name for s in doc.styles]

    def style(name, size=None, bold=None, italic=None, align=None, caps=None, before=None,
              after=None, indent=None):
        if name not in names:
            return
        st = doc.styles[name]
        if size:
            st.font.size = Pt(size)
        if bold is not None:
            st.font.bold = bold
        if italic is not None:
            st.font.italic = italic
        if caps is not None:
            st.font.small_caps = caps
        pf = st.paragraph_format
        if align is not None:
            pf.alignment = align
        if before is not None:
            pf.space_before = Pt(before)
        if after is not None:
            pf.space_after = Pt(after)
        if indent is not None:
            pf.first_line_indent = Pt(indent)
        pf.line_spacing = 1.0

    # IEEEtran journal sizes: body 10 pt, title 24 pt, authors 11 pt, abstract and index
    # terms 9 pt bold, section heads small caps centred, subsection heads italic,
    # captions 8 pt, tables 7 pt (\scriptsize), references 8 pt.
    J, C, L = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        style(name, 10, align=J, before=0, after=0, indent=10)
    style("Title", 24, bold=False, align=C, before=0, after=6, indent=0)
    style("Author", 11, align=C, before=0, after=12, indent=0)
    style("Abstract", 9, bold=True, align=J, before=0, after=0, indent=10)
    style("Heading 1", 10, bold=False, italic=False, caps=True, align=C, before=8, after=4,
          indent=0)
    style("Heading 2", 10, bold=False, italic=True, align=L, before=6, after=3, indent=0)
    style("Heading 3", 10, bold=False, italic=True, align=L, before=3, after=0, indent=10)
    style("Image Caption", 8, italic=False, align=J, before=3, after=8, indent=0)
    style("Captioned Figure", align=C, before=6, after=0, indent=0)
    style("Footnote Text", 8, indent=10)
    if "Heading 1" in names:
        doc.styles["Heading 1"].font.all_caps = False

    paras = list(doc.paragraphs)
    for p in paras:
        if p.style.name == "Abstract Title":
            p._p.getparent().remove(p._p)
    abstract = [p for p in doc.paragraphs if p.style.name == "Abstract"]
    if abstract:
        first = abstract[0]
        run = first.add_run("Abstract—")
        run.bold = run.italic = True
        if len(first.runs) > 1:
            first.runs[0]._r.addprevious(run._r)
    for p in doc.paragraphs:
        text = p.text
        if text.startswith("Index Terms"):
            p.style = doc.styles["Abstract"]
            for r in p.runs:
                r.bold = True
            p.paragraph_format.space_before = Pt(6)
        elif re.match(r"TABLE [IVX]+\. ", text):
            # IEEE puts "TABLE I" on its own line and the caption in small caps below it.
            first = p.runs[0]
            num, rest = re.match(r"(TABLE [IVX]+)\. (.*)", first.text, re.S).groups()
            first.text = rest
            head = p.add_run(num)
            first._r.addprevious(head._r)
            head.add_break()
            p.style = doc.styles["Normal"]
            p.alignment = C
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.size = Pt(8)
                r.font.small_caps = r is not head
        elif text.startswith(ALG_MARK.split("%")[0]):
            m = re.match(r"ALGLINE(\w+)QQ", text)
            kind = m.group(1)
            strip_marker(p, m.group(0))
            p.style = doc.styles["Normal"]
            p.alignment = L
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.keep_with_next = kind != "end"
            if kind == "cap":
                p._p.get_or_add_pPr().append(rule("w:pBdr", ("top", "bottom"), 8))
                pf.space_before = Pt(8)
            elif kind == "end":
                p._p.get_or_add_pPr().append(rule("w:pBdr", ("top",), 4))
                pf.space_after = Pt(6)
                for r in p.runs:
                    r.font.size = Pt(2)
            elif kind == "in":
                pf.left_indent = Pt(0)
            else:
                # "N:" in a hanging margin, then the statement indented by its nesting depth
                pf.left_indent = Pt(16 + 12 * int(kind))
                pf.first_line_indent = Pt(-16 - 12 * int(kind))
                num = p.runs[0]
                m = re.match(r"(\d+):(?: (.*))?$", num.text, re.S)
                if m:
                    num.text = m.group(2) or ""
                    lead = p.add_run(m.group(1) + ":\t")
                    num._r.addprevious(lead._r)
                    lead.font.size = Pt(8)
                    pf.tab_stops.add_tab_stop(Pt(16 + 12 * int(kind)))
                    # pandoc puts the space after "N:" in a run of its own before math
                    runs = p.runs
                    if not num.text and len(runs) > 2 and not runs[2].text.strip():
                        runs[2].text = ""
        elif p.style.name == "Captioned Figure":
            p.paragraph_format.keep_with_next = True
    for t in doc.tables:
        booktabs(t)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.style = doc.styles["Normal"]
                    p.alignment = L
                    p.paragraph_format.first_line_indent = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(7)

    # References from build/main.bbl, in IEEEtran's own wording.
    mark = next(p for p in doc.paragraphs if p.text.strip() == REF_MARK)
    head = mark.insert_paragraph_before("References", style="Heading 1")
    head.alignment = C
    entries = bbl_entries()
    if len(entries) != len(order):
        sys.exit("main.bbl has %d entries, expected %d" % (len(entries), len(order)))
    for i, text in enumerate(entries, 1):
        p = mark.insert_paragraph_before()
        p.style = doc.styles["Normal"]
        p.alignment = J
        add_runs(p, "[%d]\t%s" % (i, text))
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = Pt(18), Pt(-18)
        pf.tab_stops.add_tab_stop(Pt(18))
        for r in p.runs:
            r.font.size = Pt(8)
    mark._p.getparent().remove(mark._p)

    # Biographies: name in bold, 8 pt, as IEEEbiographynophoto sets them.
    for p in doc.paragraphs:
        if p.runs and p.runs[0].bold and p.text.startswith(("Dakshrajsinh Tank is",
                                                             "Badrinath N is")):
            for r in p.runs:
                r.font.size = Pt(8)

    # Page: US Letter, IEEE margins; title block in one column, body in two, and wide
    # (figure*, table*) floats across both columns.
    body = doc.sections[-1]
    body.page_width, body.page_height = Cm(21.59), Cm(27.94)
    body.top_margin, body.bottom_margin = Cm(1.905), Cm(2.54)
    body.left_margin = body.right_margin = Cm(1.5875)
    body.header_distance = Cm(1.0)
    template = copy.deepcopy(body._sectPr)
    set_columns(body._sectPr, 2)
    for p in list(doc.paragraphs):
        marker = p.text.strip()
        if marker == TITLE_END:
            section_break(p, 1, template)
        elif marker == WIDE_START:
            section_break(p, 2, template)
        elif marker == WIDE_END:
            section_break(p, 1, template)

    # Running head "IEEE ACCESS" with the page number at the right, as in the PDF.
    first = doc.sections[0]
    hp = first.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run("IEEE ACCESS\t")
    r.font.size = Pt(8)
    page_field(hp)
    doc.save(path)


def page_field(par):
    for kind in ("begin", "instr", "separate", "text", "end"):
        r = par.add_run()
        r.font.size = Pt(8)
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = "1"
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r._r.append(el)


def main(out=OUT) -> int:
    """Optional argument: another output path, e.g. while the usual one is open in Word."""
    tex, order, notes = to_latex()
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    pypandoc.convert_text(tex, "docx", format="latex", outputfile=out,
                          extra_args=["--resource-path", LATEX])
    finish(out, order, notes)
    print("wrote %s (%d references)" % (out, len(order)))
    return 0


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:]))
