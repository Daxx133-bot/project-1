from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path
TNR = "Times New Roman"

def _fr(run, size=12, bold=False, italic=False):
    run.font.name = TNR; run.font.size = Pt(size)
    run.font.bold = bold; run.font.italic = italic

def _fp(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, ls=1.15, sb=0, sa=6, fi=None, li=None):
    pf = p.paragraph_format; pf.alignment = align
    pf.line_spacing = ls; pf.space_before = Pt(sb); pf.space_after = Pt(sa)
    if fi is not None: pf.first_line_indent = Inches(fi)
    if li is not None: pf.left_indent = Inches(li)

def h1(doc, txt):
    p = doc.add_paragraph(); r = p.add_run(txt.upper())
    _fr(r, 14, bold=True); _fp(p, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 14, 8); return p

def h2(doc, txt):
    p = doc.add_paragraph(); r = p.add_run(txt)
    _fr(r, 13, bold=True); _fp(p, WD_ALIGN_PARAGRAPH.LEFT, 1.5, 10, 4); return p

def h3(doc, txt):
    p = doc.add_paragraph(); r = p.add_run(txt)
    _fr(r, 12, bold=True, italic=True); _fp(p, WD_ALIGN_PARAGRAPH.LEFT, 1.5, 6, 3); return p

def body(doc, txt, fi=0.5):
    p = doc.add_paragraph(); r = p.add_run(txt)
    _fr(r, 12); _fp(p, fi=fi, sa=6); return p

def cap(doc, txt):
    p = doc.add_paragraph(); r = p.add_run(txt)
    _fr(r, 11, italic=True); _fp(p, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 2, 8)

def eq(doc, formula, label):
    tbl = doc.add_table(rows=1, cols=2); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c0, c1 = tbl.rows[0].cells[0], tbl.rows[0].cells[1]
    c0.width = Inches(5.0); c1.width = Inches(1.0)
    p0 = c0.paragraphs[0]; p0.clear()
    r = p0.add_run(formula); _fr(r, 12, italic=True); _fp(p0, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 4, 4)
    p1 = c1.paragraphs[0]; p1.clear()
    r1 = p1.add_run(label); _fr(r1, 12); _fp(p1, WD_ALIGN_PARAGRAPH.RIGHT, 1.0, 4, 4)

def ref_entry(doc, num, txt):
    p = doc.add_paragraph()
    rn = p.add_run(f"[{num}]  "); _fr(rn, 11, bold=True)
    r = p.add_run(txt); _fr(r, 11)
    _fp(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.15, 0, 3, fi=-0.4, li=0.4)

def tbl_row(tbl, ri, vals, widths, bold_row=False, bg=None, size=10):
    row = tbl.rows[ri]
    for ci, (txt, w) in enumerate(zip(vals, widths)):
        c = row.cells[ci]; c.width = Inches(w)
        p = c.paragraphs[0]; p.clear()
        r = p.add_run(txt); _fr(r, size, bold=bold_row)
        _fp(p, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 2, 2)
        if bg:
            tc = c._tc; tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), bg); tcPr.append(shd)
        tc = c._tc; tcPr = tc.get_or_add_tcPr()
        tcB = OxmlElement("w:tcBorders")
        for side in ("top","left","bottom","right"):
            el = OxmlElement(f"w:{side}"); el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
            tcB.append(el)
        tcPr.append(tcB)

def set_page(doc):
    for s in doc.sections:
        s.page_width = Inches(8.27); s.page_height = Inches(11.69)
        s.left_margin = Inches(1.1); s.right_margin = Inches(1.1)
        s.top_margin = Inches(1.0); s.bottom_margin = Inches(1.0)

def add_footer_pagenum(doc):
    for sec in doc.sections:
        ftr = sec.footer
        para = ftr.paragraphs[0] if ftr.paragraphs else ftr.add_paragraph()
        para.clear(); _fp(para, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 0)
        run = para.add_run(); _fr(run, 10)
        fldChar1 = OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"), "begin")
        instrText = OxmlElement("w:instrText"); instrText.text = "PAGE"; instrText.set(qn("xml:space"), "preserve")
        fldChar2 = OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"), "separate")
        fldChar3 = OxmlElement("w:fldChar"); fldChar3.set(qn("w:fldCharType"), "end")
        run._r.extend([fldChar1, instrText, fldChar2, fldChar3])

def build_title(doc):
    p = doc.add_paragraph()
    r = p.add_run("Graph-Theoretic Criticality Ranking of Microservices:\nAn Empirical Validation Across Heterogeneous Cloud-Native Architectures")
    _fr(r, 16, bold=True); _fp(p, WD_ALIGN_PARAGRAPH.CENTER, 1.5, 24, 10)
    p2 = doc.add_paragraph()
    r2 = p2.add_run("[Author Names Removed for Blind Review]")
    _fr(r2, 12, italic=True); _fp(p2, WD_ALIGN_PARAGRAPH.CENTER, 1.2, 0, 4)
    p3 = doc.add_paragraph()
    r3 = p3.add_run("School of Computer Science and Engineering, Vellore Institute of Technology")
    _fr(r3, 11); _fp(p3, WD_ALIGN_PARAGRAPH.CENTER, 1.2, 0, 6)
    p4 = doc.add_paragraph()
    r4 = p4.add_run("Manuscript received [DATE]; revised [DATE]; accepted [DATE].")
    _fr(r4, 10, italic=True); _fp(p4, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 0, 16)

def build_abstract(doc):
    h1(doc, "Abstract")
    body(doc,
        "Microservice architectures have become the dominant paradigm for deploying "
        "cloud-native applications, yet the intricate dependency webs they form render "
        "failure impact analysis exceedingly difficult. When a single service degrades, "
        "cascading backpressure can propagate through dozens of downstream call chains, "
        "causing system-wide outages that current reactive monitoring approaches detect "
        "only after the damage has already occurred. This paper proposes and empirically "
        "validates a proactive methodology that applies graph-theoretic centrality analysis "
        "to runtime Service Dependency Graphs (SDGs) extracted from distributed traces, "
        "enabling quantitative pre-incident ranking of microservices by predicted failure "
        "blast radius.", fi=0.0)
    body(doc,
        "We conduct controlled chaos fault injection experiments across two architecturally "
        "distinct DeathStarBench benchmarks: a Social Network application comprising 12 "
        "interdependent microservices and a Hotel Reservation system with 8 loosely coupled "
        "services. For each application, SDGs are automatically constructed from Jaeger "
        "distributed traces and seven centrality metrics are computed: Degree, In-Degree, "
        "Out-Degree, Betweenness, Closeness, Eigenvector, and PageRank. A Fan-out Corrected "
        "Hybrid Metric integrating weighted out-degree with normalised Betweenness is "
        "introduced to address a structural deficiency in classical centrality measures for "
        "high fan-out orchestrator services. Pumba chaos engineering is used to inject "
        "container-kill (SIGKILL) and network-latency (100 ms tc netem) faults into each "
        "service across 96 controlled experiment runs.", fi=0.0)
    body(doc,
        "Spearman rank correlation analysis reveals that the Hybrid Metric achieves the "
        "strongest correlation with system recovery time in both architectures (rho = 0.897, "
        "p < 0.001 in Social Network; rho = 0.841, p = 0.009 in Hotel Reservation), "
        "demonstrating statistically significant cross-architecture generalisation. "
        "Classical composite centrality scores, by contrast, systematically misrank "
        "high fan-out orchestrators as low-risk nodes due to their low in-degree, a "
        "phenomenon we formalise as the Centrality-Induced False Negative Problem. "
        "Outlier analysis identifies compose-post-service as the archetypal false negative, "
        "holding composite structural rank 7 while inducing empirical blast-radius rank 2. "
        "These findings provide actionable guidance for Site Reliability Engineers prioritising "
        "chaos experiment targeting, circuit breaker placement, and redundancy investment.", fi=0.0)
    kw = doc.add_paragraph()
    rk = kw.add_run("Keywords -- "); _fr(rk, 12, bold=True, italic=True)
    rk2 = kw.add_run(
        "Microservices; graph centrality; chaos engineering; fault injection; "
        "distributed tracing; blast radius; Spearman correlation; service dependency graph; "
        "DeathStarBench; resilience engineering; site reliability engineering.")
    _fr(rk2, 12, italic=True); _fp(kw, fi=0.0, sb=4, sa=6)
    doc.add_page_break()
