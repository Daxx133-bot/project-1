"""
check_overlap.py -- how much of the report's wording also appears in the project's other
texts, which a plagiarism checker may hold (the paper drafts and README are public in the
share repository; the Review-2 draft was uploaded for review).

Measure: share of the report's word 8-grams found in each source. The appendix code and
the reference list are left out, because they are meant to match.

Output: report/out/overlap.txt
Run: python report/check_overlap.py
"""
from __future__ import annotations

import glob
import os
import re
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
N = 8


def words(text):
    text = re.sub(r"\\[a-zA-Z]+\*?(\{[^}]*\})?|%\[[^\]]*\]|\{\{[^}]*\}\}|\[@[^\]]*\]", " ", text)
    return re.findall(r"[a-z0-9]+", text.lower())


def grams(ws):
    return {" ".join(ws[i:i + N]) for i in range(len(ws) - N + 1)}


def docx_text(path):
    return "\n".join(p.text for p in docx.Document(path).paragraphs)


def main() -> int:
    report = docx_text(os.path.join(HERE, "out", "Project_Report.docx"))
    report = report.rsplit("REFERENCES", 1)[0]
    mine = grams(words(report))
    sources = {
        "paper LaTeX sections": glob.glob(os.path.join(REPO, "paper", "latex", "sections", "*.tex")),
        "paper Markdown drafts (public)": glob.glob(os.path.join(REPO, "paper", "draft", "*.md")),
        "README and PREREGISTRATION (public)": [os.path.join(REPO, "README.md"),
                                                os.path.join(REPO, "PREREGISTRATION.md")],
        "Review-2 draft report": [os.path.join(REPO, "BCSE497J_Microservice_Failure_Analysis_Report.docx")],
        "university templates (sample text)": glob.glob(os.path.join(HERE, "template", "*.docx")),
    }
    lines = ["report 8-grams: %d" % len(mine)]
    for name, files in sources.items():
        g = set()
        for f in files:
            if not os.path.isfile(f):
                continue
            t = docx_text(f) if f.endswith(".docx") else open(f, encoding="utf-8").read()
            g |= grams(words(t))
        shared = mine & g
        lines.append("%-38s %5.1f%% (%d shared)" % (name, 100 * len(shared) / len(mine),
                                                    len(shared)))
    out = "\n".join(lines)
    with open(os.path.join(HERE, "out", "overlap.txt"), "w", encoding="utf-8") as fh:
        fh.write(out + "\n")
    print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
