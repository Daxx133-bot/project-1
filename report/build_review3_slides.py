"""
build_review3_slides.py -- fill the university's Review-3 PowerPoint template with this
project's content. The template's own slides, layouts, fonts and footers are kept; only the
<<...>> placeholders are replaced, and pictures are added where a slide asks for a diagram.

Input   the template with a second Literature Review slide (after slide 5) and a second
        Experiments and Results slide (after slide 11) added by the pptx skill's
        add_slide.py, so the deck has 16 slides
        paper/numbers.json        every number on the slides ({name} fields below)
        paper/latex/build/main.bbl    reference entries, as in the IEEE paper
        paper/figures/*.png       architecture and result figures
Output  the path given as the second argument

Run: python report/build_review3_slides.py <template-with-16-slides.pptx> <out.pptx>
"""
from __future__ import annotations

import copy
import json
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
FIG = os.path.join(REPO, "paper", "figures")
FONT = "Times New Roman"

with open(os.path.join(REPO, "paper", "numbers.json"), encoding="utf-8") as fh:
    V = json.load(fh)["values"]

# Works cited on the slides, numbered in order of first appearance.
REFS = ["gan2019", "sigelman2010", "zimmermann2008", "albert2000", "luo2021", "wu2020",
        "heorhiadi2016", "basiri2016", "nosek2018", "benjamini1995"]


def n(key):
    return "[%d]" % (REFS.index(key) + 1)


def val(text):
    """Fill {claim_id} fields from paper/numbers.json; fail on an unknown id."""
    def sub(m):
        if m.group(1) not in V:
            sys.exit("unknown claim id " + m.group(1))
        return str(V[m.group(1)]).replace("-", "\u2212") if m.group(1).endswith("rho") else \
            str(V[m.group(1)])
    return re.sub(r"\{([a-z0-9_]+)\}", sub, text)


def bbl():
    with open(os.path.join(REPO, "paper", "latex", "build", "main.bbl"), encoding="utf-8") as fh:
        text = fh.read()
    out = {}
    for m in re.finditer(r"\\bibitem\{([^}]+)\}(.*?)(?=\\bibitem|\\end\{thebibliography\})",
                         text, re.S):
        e = re.sub(r"\\BIBentry(ALT|STD)interwordspacing", "", m.group(2))
        e = re.sub(r"\\(emph|url)\{", "{", e)
        e = e.replace("\\newblock", " ").replace("``", "\u201c").replace("''", "\u201d")
        e = e.replace("--", "\u2013").replace("~", " ")
        e = re.sub(r"\\([&_%#$])", r"\1", e).replace("{", "").replace("}", "")
        out[m.group(1)] = " ".join(e.split())
    return out


# ── text helpers: copy the template's own paragraph and run formatting ──
def set_runs(par, texts):
    """Put texts into the paragraph's existing runs in order; empty the rest."""
    runs = par.runs
    for i, r in enumerate(runs):
        r.text = texts[i] if i < len(texts) else ""


def fill(shape, items, size=None):
    """Replace the shape's paragraphs with items. An item is a string (a bullet; **x** is
    bold) or ("head", text) for an unbulleted bold line. Paragraph and run properties are
    copied from the template's first paragraph."""
    body = shape.text_frame._txBody
    paras = body.findall(qn("a:p"))
    proto_p = copy.deepcopy(paras[0])
    proto_r = copy.deepcopy(proto_p.find(qn("a:r")))
    for p in paras:
        body.remove(p)
    for item in items:
        head = isinstance(item, tuple)
        text = item[1] if head else item
        p = copy.deepcopy(proto_p)
        for child in list(p):
            if child.tag != qn("a:pPr"):
                p.remove(child)
        if head:
            ppr = no_bullet(p)
            ppr.set("marL", "0")
            ppr.set("indent", "0")
        for k, seg in enumerate(re.split(r"\*\*", val(text))):
            if not seg:
                continue
            r = copy.deepcopy(proto_r)
            rpr = r.find(qn("a:rPr"))
            if head or k % 2 == 1:
                rpr.set("b", "1")
            if size:
                rpr.set("sz", str(int(size * 100)))
            r.find(qn("a:t")).text = seg
            p.append(r)
        body.append(p)


def no_bullet(p):
    """<a:buNone/> in the place the schema gives it: before tabLst, defRPr and extLst, and
    replacing any other bullet setting."""
    ppr = p.find(qn("a:pPr"))
    if ppr is None:
        ppr = p.makeelement(qn("a:pPr"), {})
        p.insert(0, ppr)
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip"):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    bu = ppr.makeelement(qn("a:buNone"), {})
    later = [c for c in ppr if c.tag in (qn("a:tabLst"), qn("a:defRPr"), qn("a:extLst"))]
    if later:
        later[0].addprevious(bu)
    else:
        ppr.append(bu)
    return ppr


def remove(shape):
    shape._element.getparent().remove(shape._element)


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def body_ph(slide):
    return next(s for s in slide.shapes if s.is_placeholder and s.placeholder_format.idx == 1)


def place(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def picture(slide, name, x, y, w=None, h=None):
    path = os.path.join(FIG, name)
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
    if w and h is None and pic.height > Inches(5.0):
        ratio = Inches(5.0) / pic.height
        pic.height, pic.width = Inches(5.0), int(pic.width * ratio)
    return pic


# ── literature table, styled like the template's sample table ──
LIT = [
    ("gan2019", "Gan et al., 2019", "Open benchmark suite of microservice applications",
     "End-to-end applications (Social Network, Hotel Reservation, ...) studied on real hardware",
     "Pro: realistic, open test bed. Con: synthetic workloads",
     "Performance is sensitive to network and tail latency; slowdowns spread across services"),
    ("sigelman2010", "Sigelman et al., 2010 (Dapper)",
     "Trace requests across a large distributed system",
     "Spans with parent links, sampling, always-on collection",
     "Pro: gives the caller-to-callee edges used here. Con: does not say which service matters most",
     "Low-overhead tracing is practical at production scale"),
    ("zimmermann2008", "Zimmermann and Nagappan, 2008",
     "Predict defect-prone components from dependency graphs",
     "Network measures on the dependency graph of a large commercial system",
     "Pro: graph position carries risk information. Con: outcome is defects, not run-time failure",
     "Network measures found defect-prone binaries better than complexity metrics"),
    ("albert2000", "Albert et al., 2000", "Error and attack tolerance of complex networks",
     "Node removal on model and real networks",
     "Pro: links structure and robustness. Con: failure is node removal, not latency and errors",
     "Scale-free networks survive random failures but break when hubs are removed"),
    ("luo2021", "Luo et al., 2021", "Characterise production microservice call graphs",
     "Analysis of large-scale Alibaba production traces",
     "Pro: real call graphs at scale. Con: no fault injection",
     "Call graphs are highly dynamic across requests and over time"),
    ("wu2020", "Wu et al., 2020 (MicroRCA)", "Locate root causes of performance issues",
     "Attributed graph of the running system ranked with personalised PageRank",
     "Pro: combines centrality with run-time metrics. Con: needs an anomaly first",
     "Graph-based ranking located root causes in a benchmark microservice system"),
    ("heorhiadi2016", "Heorhiadi et al., 2016 (Gremlin)",
     "Test the resilience of microservice applications",
     "Fault injection into the messages between services",
     "Pro: service-level faults. Con: checks fallbacks, not impact ranking",
     "Systematic fault injection exposes missing or faulty failure handling"),
    ("basiri2016", "Basiri et al., 2016", "Chaos engineering at Netflix",
     "Controlled fault experiments on production systems",
     "Pro: faults used to learn how a system fails. Con: aims at confidence, not ranking",
     "Live experiments build confidence that a system tolerates failure"),
]


def lit_table(slide, rows):
    head = ("Paper", "Objective", "Methodology", "Pros/Cons", "Findings")
    widths = (2.2, 2.1, 2.4, 2.4, 2.1)
    t = slide.shapes.add_table(len(rows) + 1, 5, Inches(0.87), Inches(1.55),
                               Inches(sum(widths)), Inches(4.9)).table
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
    fills = (RGBColor(0x4F, 0x81, 0xBD), RGBColor(0xD0, 0xD8, 0xE8), RGBColor(0xE9, 0xED, 0xF4))
    for i, row in enumerate([head] + [(r[1] + " " + n(r[0]),) + r[2:] for r in rows]):
        for j, text in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = fills[0] if i == 0 else fills[1 + (i % 2 == 0)]
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            p = c.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = text
            r.font.name, r.font.size = FONT, Pt(16 if i == 0 else 11)
            r.font.bold = i == 0
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else RGBColor(0, 0, 0)


def main(src, out):
    prs = Presentation(src)
    s = prs.slides
    refs = bbl()

    # 1 title
    title = shape(s[0], "Title 6").text_frame.paragraphs[0]
    set_runs(title, ["B.Tech", "", "", "", "", "BCSE497J \u2013 Project-I", "",
                     "Identifying Failure-Prone Microservices Using Graph Centrality Metrics"])
    for r in title.runs:                       # the template highlights its fill-in fields
        for mark in r._r.findall(".//" + qn("a:highlight")):
            mark.getparent().remove(mark)
    team = shape(s[0], "object 3").text_frame.paragraphs
    set_runs(team[2], ["Dakshrajsinh Tank\t\t  (", "23BDS0299", ")"])
    set_runs(team[6], ["", "Dr. Badrinath N", ""])
    team[3]._p.getparent().remove(team[3]._p)                 # one-member team

    # 2 approval mail: left for the student's own approval mail

    # 3 aim
    fill(body_ph(s[2]), [
        "To test, with a fault-injection experiment planned and locked in advance, whether "
        "the position of a service in a microservice call graph predicts how far its failure "
        "spreads",
        "Two graph measures are compared: the number of callers of a service, and a "
        "fan-out-weighted centrality metric that ranks services calling many others as the "
        "riskiest"], size=24)

    # 4 abstract
    fill(body_ph(s[3]), [
        "**Problem:** choosing which microservices to harden first; a ranking from the call "
        "graph is cheap, injecting faults into every service is not",
        "**Method:** {campaign_runs_total} fault-injection runs on {campaign_architectures} "
        "DeathStarBench applications " + n("gan2019") + "; every service received a process "
        "kill and a {fault_latency_delay_ms} ms network delay, {campaign_reps_per_cell} times "
        "each; graphs rebuilt from traces; latency and errors measured on each service's own "
        "spans",
        "**Result:** caller count and degraded callers correlated at \u03c1 = {primary_rho} "
        "(Social Network, kill), but no test survived correction (adjusted p = "
        "{primary_p_bh_locked}); the fan-out metric was negative everywhere and never "
        "distinguishable from zero",
        "**Clearest finding:** no callee of the faulted service degraded in "
        "{descendant_affected_runs_zero} of {campaign_runs_total} runs \u2013 failures travel "
        "to callers"], size=18)

    # 5 and 6 literature review (two slides, as the template allows)
    for slide, rows in ((s[4], LIT[:4]), (s[5], LIT[4:])):
        for sh in list(slide.shapes):
            if sh.name in ("Content Placeholder 4", "Picture 6"):   # sample table, note
                remove(sh)
        lit_table(slide, rows)
    shape(s[5], "object 2").text_frame.paragraphs[0].runs[0].text = "Literature Review (contd.)"
    for r in shape(s[5], "object 2").text_frame.paragraphs[0].runs[1:]:
        r.text = ""

    # 7 research gap
    fill(body_ph(s[6]), [
        "Graph measures have ranked software components by defect risk " + n("zimmermann2008")
        + ", and network studies model failure as removing a node " + n("albert2000"),
        "Fault-injection tools test whether a system survives failures " + n("heorhiadi2016")
        + ", " + n("basiri2016") + "; root-cause methods use the graph only after an anomaly "
        "has started " + n("wu2020"),
        "No work found tests, with a plan fixed in advance, whether a graph-based ranking "
        "predicts how far failures spread in a running microservice system",
        "The fan-out intuition and the opposite, caller-based argument had not been compared "
        "against measured outcomes"], size=20)

    # 8 objectives, SDG, TRL, outcome
    ph = body_ph(s[7])
    fill(ph, [
        "Rebuild each application's call graph from traces; compute caller count, the "
        "fan-out metric and comparator measures",
        "Lock hypotheses and analysis before collecting data (preregistration "
        + n("nosek2018") + ")",
        "Run a balanced fault-injection campaign and measure every service on its own spans",
        "Test which graph measure tracks the number of degraded callers, with correction "
        "for multiple tests"], size=18)
    place(ph, 0.92, 1.55, 11.5, 2.85)
    box = shape(s[7], "TextBox 6")
    paras = box.text_frame.paragraphs
    set_runs(paras[0], ["Identified:  SDG 9 \u2013 Industry, Innovation and Infrastructure   |   "
                        "TRL 3 \u2013 Experimental proof of concept   |   ",
                        "Outcome: ", "Journal \u2013 IEEE Access (manuscript prepared)"])
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    place(box, 0.5, 4.4, 12.3, 0.45)

    # 9 architecture
    ph = body_ph(s[8])
    picture(s[8], "fig_pipeline.png", 0.6, 1.6, h=5.2)
    fill(ph, [
        "Graph reconstruction from Jaeger traces",
        "Fault injection with Pumba: kill ({fault_kill_signal}) or "
        "{fault_latency_delay_ms} ms delay",
        "Per-service latency and errors from each service's own spans",
        "Degraded services classified as callers, callees or unrelated",
        "Correlation tests with Benjamini-Hochberg correction " + n("benjamini1995")],
        size=16)
    place(ph, 8.1, 1.7, 4.6, 5.0)

    # 10 functional requirements
    fill(body_ph(s[9]), [
        "Collect traces from Jaeger and rebuild a caller-to-callee graph for each application",
        "Compute caller count, the fan-out metric and comparator centralities for every service",
        "Find the target container and inject a kill or a {fault_latency_delay_ms} ms delay; "
        "restore the service afterwards",
        "Measure p95 latency and error rate of every service before and during the fault",
        "Classify degraded services by their relation to the faulted one and measure "
        "recovery time",
        "Check run integrity, run the statistics, and regenerate every number by script"],
        size=18)

    # 11 modules
    fill(body_ph(s[10]), [
        "**Graph reconstruction:** tools/extract_graph.py",
        "**Predictor computation:** centrality/compute_centrality.py, "
        "tools/build_predictor_table.py",
        "**Campaign runner and fault injection:** tools/run_campaign.py, tools/run_pilot.py",
        "**Measurement:** measurement/metrics_collector.py, blast_radius.py, recovery_probe.py",
        "**Analysis and reporting:** analysis/final/run_confirmatory_analysis.py, "
        "paper/build_all.py"], size=18)

    # 12 experiments and results (1): set-up and the caller count
    ph = body_ph(s[11])
    fill(ph, [
        "{sn_services} Social Network and {hr_services} Hotel Reservation services on Docker; "
        "{campaign_runs_total} runs in {campaign_cells} balanced cells",
        "Caller count vs degraded callers (Social Network, kill): \u03c1 = {primary_rho}, "
        "raw p = {primary_p_locked}",
        "After correction across {conf_family_size} tests " + n("benjamini1995")
        + ": adjusted p = {primary_p_bh_locked}, so the hypothesis is **not confirmed**"],
        size=18)
    place(ph, 0.92, 1.55, 11.5, 1.9)
    picture(s[11], "fig_primary_sn.png", 2.55, 3.55, h=3.3)

    # 13 experiments and results (2): fan-out metric and direction
    ph = body_ph(s[12])
    shape(s[12], "object 2").text_frame.paragraphs[0].runs[0].text = \
        "Experiments and Results (contd.)"
    fill(ph, [
        "Fan-out metric vs degraded callers: \u03c1 = {h1b_sn_kill_rho} (Social Network) and "
        "{h1b_hr_kill_rho} (Hotel Reservation); never distinguishable from zero",
        "None of {ws_triplets} weightings of the metric gave a significant result",
        "No callee of the faulted service degraded in {descendant_affected_runs_zero} of "
        "{campaign_runs_total} runs: failures travel to callers",
        "Charging end-to-end latency to every service on a path flags unrelated services; "
        "per-service measurement avoids this"], size=17)
    place(ph, 0.92, 1.55, 6.2, 5.2)
    picture(s[12], "fig_rank_inversion.png", 7.3, 1.75, w=5.4)

    # 14 conclusion and future work
    fill(body_ph(s[13]), [
        ("head", "Conclusion"),
        "The fan-out-weighted centrality metric did not identify failure-prone services",
        "Caller count correlated with degraded callers, but not significantly after "
        "correction, and partly by construction",
        "Failures travelled to callers and almost never to callees",
        ("head", "Future work"),
        "Test a traffic-weighted caller ranking as a new preregistered hypothesis",
        "Use larger applications to raise statistical power",
        "Study how retries and circuit breakers change failure spread"], size=18)

    # 15 outcome achieved
    fill(body_ph(s[14]), [
        "**Journal manuscript (IEEE Access, Scopus indexed):** \u201cDoes Call-Graph Fan-Out "
        "Predict Failure Impact in Microservices? A Preregistered Fault-Injection Study\u201d",
        "**Status:** manuscript prepared; submission pending",
        "**Code and data:** public repository that regenerates every reported number "
        "(github.com/Daxx133-bot/proj-1-updated)"], size=20)

    # 16 references, in the paper's IEEE format
    ph = shape(s[15], "Rectangle 3")
    fill(ph, ["[%d] %s" % (i, refs[k]) for i, k in enumerate(REFS, 1)], size=11)
    place(ph, 0.6, 1.5, 11.9, 5.3)
    for p in ph.text_frame._txBody.findall(qn("a:p")):
        no_bullet(p)
        for r in p.findall(qn("a:r")):
            r.find(qn("a:rPr")).attrib.pop("b", None)

    prs.save(out)
    print("wrote %s (%d slides)" % (out, len(prs.slides)))
    return 0


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:3]))
